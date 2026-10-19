#!/usr/bin/env python3
"""Convert the personal-agent PRD markdown into a Word document for review."""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("PRD-个人桌面Agent.md")
DST = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("PRD-个人桌面Agent.docx")

CN_FONT = "微软雅黑" if sys.platform == "win32" else "PingFang SC"
MONO_FONT = "Menlo" if sys.platform == "darwin" else "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x59, 0x59, 0x59)
CODE_SHADE = "F4F5F7"
HEADER_SHADE = "E8EEF6"


def set_run_font(run, name):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def shade(element, fill):
    pr = element.get_or_add_tcPr() if element.tag.endswith("tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def add_inline(paragraph, text, base_bold=False):
    """Handle **bold**, `code`, and plain text within one paragraph."""
    for token in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, MONO_FONT)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        else:
            run = paragraph.add_run(token)
        set_run_font(run, CN_FONT)
        if base_bold:
            run.bold = True


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_sep(line):
    return set(line.strip()) <= set("|-: ") and "-" in line


def add_table(doc, rows):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    usable = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    for i, cells in enumerate(rows):
        row = table.add_row()
        for j in range(cols):
            cell = row.cells[j]
            cell.width = int(usable / cols)
            text = cells[j] if j < len(cells) else ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            add_inline(para, text, base_bold=(i == 0))
            if i == 0:
                shade(cell._tc, HEADER_SHADE)
            for run in para.runs:
                run.font.size = Pt(9)
    doc.add_paragraph()


def build(md_text):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for attr in ("top_margin", "bottom_margin"):
        setattr(section, attr, Cm(1.9))
    for attr in ("left_margin", "right_margin"):
        setattr(section, attr, Cm(2.0))

    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = CN_FONT
    rpr = normal.element.get_or_add_rPr()
    rfonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(attr), CN_FONT)
    rpr.append(rfonts)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.25

    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.color.rgb = ACCENT
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(4)

    lines = md_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.4)
            para.paragraph_format.space_before = Pt(4)
            para.paragraph_format.space_after = Pt(8)
            shade(para._p, CODE_SHADE)
            run = para.add_run("\n".join(block))
            set_run_font(run, MONO_FONT)
            run.font.size = Pt(8.5)
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_sep(lines[i + 1]):
            rows = [split_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        if stripped.startswith("---") and set(stripped) <= set("-"):
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            ppr = para._p.get_or_add_pPr()
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "BFBFBF")
            border.append(bottom)
            ppr.append(border)
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip().replace("**", "")
            if level == 1:
                para = doc.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run(text)
                set_run_font(run, CN_FONT)
                run.font.size = Pt(20)
                run.font.bold = True
                run.font.color.rgb = ACCENT
                para.paragraph_format.space_after = Pt(10)
            else:
                doc.add_heading(text, level=min(level - 1, 3))
            i += 1
            continue

        if stripped.startswith("> "):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.6)
            shade(para._p, "F7F7F2")
            add_inline(para, stripped[2:])
            for run in para.runs:
                run.font.color.rgb = MUTED
                run.font.size = Pt(10)
            i += 1
            continue

        if re.match(r"^[-*] \[[ x]\] ", stripped):
            checked = stripped[3] == "x"
            para = doc.add_paragraph(style="List Bullet")
            add_inline(para, ("☑ " if checked else "☐ ") + stripped[6:])
            i += 1
            continue

        if re.match(r"^[-*] ", stripped):
            para = doc.add_paragraph(style="List Bullet")
            add_inline(para, stripped[2:])
            i += 1
            continue

        if re.match(r"^\d+\. ", stripped):
            para = doc.add_paragraph(style="List Number")
            add_inline(para, re.sub(r"^\d+\. ", "", stripped))
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        para = doc.add_paragraph()
        add_inline(para, stripped)
        i += 1

    doc.save(DST)
    print(f"wrote {DST}")


build(SRC.read_text(encoding="utf-8"))
